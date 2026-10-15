#!/usr/bin/env python3
from __future__ import annotations

import json, re, hashlib
from pathlib import Path
import numpy as np
import pandas as pd
from scipy import stats
from sklearn.metrics import r2_score, mean_squared_error, mean_absolute_error
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

ROOT=Path(__file__).resolve().parents[2]
FIG5=ROOT/'fig5'
OUT=ROOT/'supplementary_tables'
OUT.mkdir(exist_ok=True)

FEATURES=[
('−NCP','Opposite_Charge_Number','bulk electrostatics','pair','−(net charge of RGG partner × net charge of TAD partner). Historical internal column Opposite_Charge_Number is exactly −NCP.','Bulk charge complementarity between sequence partners.'),
('mean λ RGG','mean_lambda_RGG','hydropathy/stickiness','RGG partner','lambda_RGG divided by RGG_length.','Mean CALVADOS hydropathy/stickiness of the RGG partner.'),
('mean λ TAD','mean_lambda_TAD','hydropathy/stickiness','TAD partner','lambda_TAD divided by TAD_length.','Mean CALVADOS hydropathy/stickiness of the TAD partner.'),
('SHD RGG','SHD_RGG','sequence patterning','RGG partner','Sequence hydropathy decoration descriptor computed for the RGG partner.','Hydropathy patterning along the RGG sequence.'),
('SHD TAD','SHD_TAD','sequence patterning','TAD partner','Sequence hydropathy decoration descriptor computed for the TAD partner.','Hydropathy patterning along the TAD sequence.'),
('local charge-complementarity heterogeneity','opp_charge_std','local electrostatics','overlapping 5-aa patches across both partners','Population standard deviation of local opposite-charge/complementarity scores across validated historical overlapping 5-aa patch pairs.','Local heterogeneity in electrostatic complementarity beyond the bulk −NCP descriptor.'),
('SCD patch 1 heterogeneity','scd_patch1_std','local charge patterning','RGG overlapping 5-aa patches','Population standard deviation of SCD over RGG-side 5-aa patches from the validated historical patch implementation.','Local charge-pattern heterogeneity along the RGG partner.'),
('SCD patch 2 heterogeneity','scd_patch2_std','local charge patterning','TAD overlapping 5-aa patches','Population standard deviation of SCD over TAD-side 5-aa patches from the validated historical patch implementation.','Local charge-pattern heterogeneity along the TAD partner.'),
('RG fraction patch 1 heterogeneity','rg_fraction_patch1_std','local composition','RGG overlapping 5-aa patches','Population standard deviation of R/G residue fraction across RGG-side 5-aa patches.','Local R/G compositional heterogeneity of the RGG partner.'),
('aromatic fraction patch 2 heterogeneity','aromatic_fraction_patch2_std','local composition','TAD overlapping 5-aa patches','Population standard deviation of aromatic fraction across TAD-side 5-aa patches.','Local aromatic compositional heterogeneity of the TAD partner.'),
('SHD patch 1 heterogeneity','shd_patch1_std','local hydropathy patterning','RGG overlapping 5-aa patches','Population standard deviation of patch-level SHD across RGG-side 5-aa patches.','Local hydropathy-pattern heterogeneity along the RGG partner.'),
('SHD patch 2 heterogeneity','shd_patch2_std','local hydropathy patterning','TAD overlapping 5-aa patches','Population standard deviation of patch-level SHD across TAD-side 5-aa patches.','Local hydropathy-pattern heterogeneity along the TAD partner.'),
('AH energy heterogeneity','ah_energy_std','local pair interaction','all overlapping 5-aa RGG × TAD patch pairs','Population standard deviation of Ashbaugh-Hatch interaction scores across validated historical 5-aa patch-pair implementation.','Heterogeneity in local short-range hydropathy/stickiness interaction potential.'),
]
FEATURE_ORDER=[x[1] for x in FEATURES]

def sha(p):
    h=hashlib.sha256();
    with open(p,'rb') as f:
        for b in iter(lambda:f.read(1024*1024),b''): h.update(b)
    return h.hexdigest()

def metrics(y,p):
    y=np.asarray(y,float); p=np.asarray(p,float); ok=np.isfinite(y)&np.isfinite(p); y=y[ok]; p=p[ok]
    n=len(y); resid=y-p
    return {'n':n,'R2':r2_score(y,p) if n>1 else np.nan,'RMSE':mean_squared_error(y,p)**0.5 if n else np.nan,'MAE':mean_absolute_error(y,p) if n else np.nan,'Pearson':stats.pearsonr(y,p).statistic if n>2 else np.nan,'Spearman':stats.spearmanr(y,p).statistic if n>2 else np.nan,'bias_sim_minus_pred':float(np.mean(resid)) if n else np.nan}

def corr(x,y):
    x=np.asarray(x,float); y=np.asarray(y,float); ok=np.isfinite(x)&np.isfinite(y); x=x[ok]; y=y[ok]; n=len(x)
    return {'n':n,'Pearson':stats.pearsonr(x,y).statistic if n>2 else np.nan,'Spearman':stats.spearmanr(x,y).statistic if n>2 else np.nan}

def parse_audit(path):
    txt=Path(path).read_text() if Path(path).exists() else ''
    vals={}
    for key in ['lambda_loaded','cap_loaded','gate_center','gate_scale']:
        m=re.search(rf'{key}=([^\n]+)',txt); vals[key]=m.group(1) if m else ''
    m=re.search(r'development_identity_diffs=([^\n]+)',txt); vals['development_identity_diffs']=m.group(1) if m else ''
    m=re.search(r'model_artifact_hashes=([^\n]+)',txt); vals['model_artifact_hashes']=m.group(1) if m else ''
    return vals,txt

def clean_df(df):
    df=df.copy()
    ren={'pair_id':'pair ID','name':'pair ID','minus_NCP':'−NCP','Opposite_Charge_Number':'−NCP / Opposite_Charge_Number','simulated_target':'simulated log10(−B22)','B22_corr1_log10':'simulated log10(−B22)','base_prediction':'base prediction','final_hybrid_prediction':'final hybrid prediction','OOF_base_prediction':'OOF base prediction','OOF_final_hybrid_prediction':'OOF final hybrid prediction','OOF_applied_correction':'OOF applied residual correction','applied_correction':'applied residual correction','raw_residual_prediction':'raw residual prediction','gate_value':'gate value','lambda_final':'lambda','absolute_error':'absolute error','high_minus_NCP':'high −NCP flag','is_rbp':'RBP status','residual':'residual','target_evaluable':'target evaluable','Corrected_B22_Hummer':'raw B22','simulated_corrected_B22':'raw B22','parent_WT_pair_id':'WT parent pair ID'}
    return df.rename(columns={k:v for k,v in ren.items() if k in df.columns})

def write_df(ws, df, start_row=1, start_col=1):
    df=df.copy()
    for col in df.columns:
        if df[col].dtype==object:
            df[col]=df[col].where(pd.notna(df[col]), '')
    for j,c in enumerate(df.columns,start_col):
        cell=ws.cell(start_row,j,c); cell.font=Font(bold=True,color='FFFFFF'); cell.fill=PatternFill('solid',fgColor='4F81BD'); cell.alignment=Alignment(wrap_text=True,vertical='top')
    for i,row in enumerate(df.itertuples(index=False),start_row+1):
        for j,val in enumerate(row,start_col):
            if pd.isna(val): val=None
            ws.cell(i,j,val)
    max_row=start_row+len(df); max_col=start_col+len(df.columns)-1
    if len(df.columns)>0:
        ws.auto_filter.ref=f"{get_column_letter(start_col)}{start_row}:{get_column_letter(max_col)}{max_row}"
    ws.freeze_panes=ws.cell(start_row+1,start_col).coordinate
    for j,c in enumerate(df.columns,start_col):
        sample=df[c].head(100).fillna('').tolist()
        max_len=max([len(str(c))]+[len(str(x)) for x in sample])
        ws.column_dimensions[get_column_letter(j)].width=min(max(max_len+2,10),45)
    return max_row,max_col

def add_readme(wb, lines):
    ws=wb.active; ws.title='README'
    ws['A1']='README'; ws['A1'].font=Font(bold=True,size=14)
    r=3
    for k,v in lines:
        ws.cell(r,1,k).font=Font(bold=True)
        ws.cell(r,2,str(v)); ws.cell(r,2).alignment=Alignment(wrap_text=True,vertical='top')
        r+=1
    ws.column_dimensions['A'].width=30; ws.column_dimensions['B'].width=120; ws.freeze_panes='A3'

def add_sheet(wb,name,df):
    ws=wb.create_sheet(name[:31]); write_df(ws,df); return ws

def save_wb(wb,path):
    for ws in wb.worksheets: ws.sheet_view.showGridLines=True
    wb.save(path)

fig5_final=FIG5/'outputs/final_hybrid_model/final_predictions'
controls_pred=FIG5/'controls/final_predictions'
finches=ROOT/'supplementary/finches/figS6_finches_scaled'
cluster_vals,cluster_text=parse_audit(controls_pred/'FINAL_CONTROL_PREDICTION_AUDIT.txt')

# S6
wb=Workbook(); add_readme(wb,[('Workbook','Table S6. Final hybrid predictor resources for Figure 5 and Figure S5.'),('Target','log10(−B22), defined where B22 is negative/attractive.'),('Development/test design','Model development used 917 simulated pairs. The 1,794 sequence-exclusive pairs are the locked test set. The full predicted interaction landscape contains definitive frozen-model predictions.'),('Model summary','The final hybrid model combines global sequence descriptors and validated historical local-context descriptors. The residual NN is a fixed correction component.'),('−NCP definition','−NCP is the negative net-charge product. The historical internal feature Opposite_Charge_Number is exactly −NCP.'),('Frozen deployment provenance',f"lambda={cluster_vals.get('lambda_loaded')}; cap={cluster_vals.get('cap_loaded')}; gate centre={cluster_vals.get('gate_center')}; gate scale={cluster_vals.get('gate_scale')}; hashes={cluster_vals.get('model_artifact_hashes')}"),('Identity check',cluster_vals.get('development_identity_diffs')),('Omissions','Old B22_pred_hybrid, lambda=0 outputs, preliminary candidate predictions, and obsolete Normalized OCN terminology are excluded.')])
add_sheet(wb,'Feature dictionary',pd.DataFrame(FEATURES,columns=['publication feature name','internal column name','feature family','sequence/partner scope','mathematical/computational definition','biological/computational interpretation']))
fam_path=FIG5/'reviewer_feature_family_decomposition_training_only/cv_metrics_summary.csv'
if fam_path.exists():
    fam=pd.read_csv(fam_path); fam_p=fam.pivot(index='model',columns='metric',values='mean').reset_index(); fam_p['n']=917; fam_p['CV scheme']='5-fold repeated training-only CV; held-out set not used for feature selection'; fam_p['included features/families']='see model label; reviewer feature-family decomposition'; fam_p=fam_p.rename(columns={'model':'model label','r2':'R²','rmse':'RMSE','mae':'MAE','pearson_r':'Pearson','spearman_rho':'Spearman'})
else: fam_p=pd.DataFrame()
add_sheet(wb,'Fig5B model development',fam_p)
by=FIG5/'reviewer_f2_patch_context_training_only/cv_metrics_by_repeat.csv'
if by.exists():
    bd=pd.read_csv(by); patch_sum=bd.groupby('model').agg(n=('repeat_seed','count'),R2=('R2','mean'),RMSE=('RMSE','mean'),MAE=('MAE','mean'),Pearson=('Pearson','mean'),Spearman=('Spearman','mean')).reset_index(); base=float(patch_sum.loc[patch_sum['model'].eq('F2'),'R2'].iloc[0]); patch_sum['ΔR² relative to global base']=patch_sum['R2']-base; desc={'F2':'global base: −NCP, mean λ RGG/TAD, SHD RGG/TAD','F2-P1':'F2 + local charge-complementarity heterogeneity','F2-P7':'F2 + prespecified balanced local-context panel','F2-P8':'F2-P7 + AH energy heterogeneity; final feature set'}; patch_sum.insert(1,'publication label',patch_sum['model'].map(desc).fillna(''))
else: patch_sum=pd.DataFrame()
add_sheet(wb,'Fig5C patch ablation',patch_sum)
for sheet,path,cols in [('Development OOF',fig5_final/'development_917_OOF_final_hybrid.csv',['pair_id','simulated_target','minus_NCP','OOF_base_prediction','OOF_final_hybrid_prediction','OOF_applied_correction','residual','absolute_error','high_minus_NCP','model_manifest_sha256']),('Locked test 1794',fig5_final/'test_1794_final_hybrid_predictions.csv',['pair_id','TAD_id','RGG_id','simulated_target','minus_NCP','base_prediction','raw_residual_prediction','gate_value','applied_correction','final_hybrid_prediction','residual','absolute_error','high_minus_NCP','is_rbp','model_manifest_sha256']),('Full landscape',fig5_final/'full_landscape_final_hybrid_predictions.csv',['pair_id','TAD_id','RGG_id','minus_NCP','base_prediction','applied_correction','final_hybrid_prediction','is_rbp','high_minus_NCP','model_manifest_sha256'])]:
    df=pd.read_csv(path,usecols=lambda c:c in cols); df=df[[c for c in cols if c in df.columns]]; add_sheet(wb,sheet,clean_df(df))
res=pd.read_csv(fig5_final/'development_917_OOF_final_hybrid.csv')
res_src=pd.DataFrame({'pair ID':res['pair_id'],'simulated log10(−B22)':res['simulated_target'],'−NCP':res['minus_NCP'],'base OOF prediction':res['OOF_base_prediction'],'base residual':res['simulated_target']-res['OOF_base_prediction'],'hybrid OOF prediction':res['OOF_final_hybrid_prediction'],'hybrid residual':res['residual'],'applied residual correction':res['OOF_applied_correction'],'high −NCP flag':res['high_minus_NCP']})
add_sheet(wb,'FigS5 residual source',res_src)
add_sheet(wb,'Residual summary',pd.DataFrame([{'scope':'all development pairs','model':'base OOF',**metrics(res['simulated_target'],res['OOF_base_prediction'])},{'scope':'all development pairs','model':'hybrid OOF',**metrics(res['simulated_target'],res['OOF_final_hybrid_prediction'])}]))
add_sheet(wb,'Deployment provenance',pd.DataFrame([{'item':'development n','value':917},{'item':'locked test n','value':1794},{'item':'final feature order','value':'; '.join(FEATURE_ORDER)},{'item':'base-model architecture','value':'StandardScaler plus MLP ensemble from frozen model package'},{'item':'residual-network architecture','value':'fixed residual neural-network correction component from frozen model package'},{'item':'lambda','value':cluster_vals.get('lambda_loaded')},{'item':'residual cap','value':cluster_vals.get('cap_loaded')},{'item':'gate centre','value':cluster_vals.get('gate_center')},{'item':'gate scale','value':cluster_vals.get('gate_scale')},{'item':'target definition','value':'log10(−B22) where B22 is negative/attractive'},{'item':'cluster identity check','value':cluster_vals.get('development_identity_diffs')},{'item':'model hashes','value':cluster_vals.get('model_artifact_hashes')}]))
path_s6=OUT/'Table_S6_Final_Hybrid_Predictor.xlsx'; save_wb(wb,path_s6)

# S7 FINCHES
wb=Workbook(); audit_fin=(finches/'FINCHES_final_model_audit.txt').read_text(); add_readme(wb,[('Workbook','Table S7. FINCHES resources for Figure S6 only.'),('Scope','FINCHES is separate from frozen-model control datasets.'),('FINCHES source','FINCHES epsilon values are from the pre-existing FINCHES run and were not recomputed.'),('Figure S6A','FINCHES epsilon versus CALVADOS simulations for 917 development pairs.'),('Figure S6B','FINCHES epsilon versus final predicted full landscape.'),('Scaling','Figure scaling is visual 0–1 transformation. Raw epsilon and interaction values are retained.'),('Omissions','Historical B22_pred_hybrid and model-control data are not included.')])
fa=pd.read_csv(finches/'FigS6A_training_scaled_merged.csv'); fb=pd.read_csv(finches/'FigS6B_unsim_scaled_merged.csv')
add_sheet(wb,'S6A FINCHES simulation',fa.rename(columns={'name':'pair ID','epsilon':'raw FINCHES epsilon','eps_scaled':'scaled FINCHES value','B22_corr1_log10':'simulated interaction value','y_scaled':'scaled interaction value'}))
add_sheet(wb,'S6B FINCHES landscape',fb.rename(columns={'name':'pair ID','epsilon':'raw FINCHES epsilon','eps_scaled':'scaled FINCHES value','final_hybrid_prediction':'final hybrid prediction','y_scaled':'scaled final prediction'}))
summary_rows=[{'analysis':'S6A FINCHES vs 917 simulation','n':len(fa),'Pearson':stats.pearsonr(fa['epsilon'],fa['B22_corr1_log10']).statistic,'Spearman':stats.spearmanr(fa['epsilon'],fa['B22_corr1_log10']).statistic,'raw epsilon min':fa['epsilon'].min(),'raw epsilon max':fa['epsilon'].max(),'raw y min':fa['B22_corr1_log10'].min(),'raw y max':fa['B22_corr1_log10'].max()},{'analysis':'S6B FINCHES vs final full landscape','n':len(fb),'Pearson':stats.pearsonr(fb['epsilon'],fb['final_hybrid_prediction']).statistic,'Spearman':stats.spearmanr(fb['epsilon'],fb['final_hybrid_prediction']).statistic,'raw epsilon min':fb['epsilon'].min(),'raw epsilon max':fb['epsilon'].max(),'raw y min':fb['final_hybrid_prediction'].min(),'raw y max':fb['final_hybrid_prediction'].max()}]
ft=finches/'FINCHES_test_1794_audit.csv'
if ft.exists():
    ft_df=pd.read_csv(ft)
    if len(ft_df)>0:
        summary_rows.append({'analysis':'FINCHES vs locked 1,794 simulated test target','n':len(ft_df),'Pearson':stats.pearsonr(ft_df['epsilon'],ft_df['B22_corr1_log10']).statistic,'Spearman':stats.spearmanr(ft_df['epsilon'],ft_df['B22_corr1_log10']).statistic,'raw epsilon min':ft_df['epsilon'].min(),'raw epsilon max':ft_df['epsilon'].max(),'raw y min':ft_df['B22_corr1_log10'].min(),'raw y max':ft_df['B22_corr1_log10'].max()}); add_sheet(wb,'Test 1794 audit',ft_df.rename(columns={'name':'pair ID','epsilon':'raw FINCHES epsilon','B22_corr1_log10':'simulated test target'}))
add_sheet(wb,'Summary',pd.DataFrame(summary_rows))
orig=re.search(r'source filename: ([^\n]+)',audit_fin)
add_sheet(wb,'Provenance',pd.DataFrame([{'item':'audit provenance','value':audit_fin},{'item':'original FINCHES score file','value':orig.group(1) if orig else ''},{'item':'epsilon column','value':'epsilon'},{'item':'pair-ID mapping','value':'constructed as rgg + __ + tad from FINCHES matrix labels'},{'item':'FINCHES rerun','value':'No; values were not recomputed'}]))
path_s7=OUT/'Table_S7_FINCHES.xlsx'; save_wb(wb,path_s7)

# S8 Controls
wb=Workbook(); figs7_audit=(FIG5/'controls/figS7_final_model_controls/FIGS7_FINAL_AUDIT.txt').read_text(); add_readme(wb,[('Workbook','Table S8. Final frozen-model control and applicability resources for Figure S7.'),('Scope','Frozen-model generalization/applicability controls generated after model freezing.'),('Model-use restriction','No control dataset was used for model fitting, feature selection, or recalibrating lambda/cap/gate.'),('Prediction source','Predictions came from authoritative cluster-deployed frozen hybrid model.'),('Analyses','Non-basic WT external/control pairs; WT-to-shuffle perturbation prediction; histidine pH6.5 applicability; noncanonical/basic applicability.'),('Omissions','WT simulated vs shuffle simulated, pH7 vs pH6.5 simulated, and Figure 4 RDF/convergence data are excluded.')])
nb=pd.read_csv(controls_pred/'nonbasic_WT_final_predictions.csv'); dev_ids=set(pd.read_csv(fig5_final/'development_917_OOF_final_hybrid.csv',usecols=['pair_id'])['pair_id']); nb['overlap with 917 development set']=nb['name'].isin(dev_ids)
add_sheet(wb,'Non-basic WT',clean_df(nb[[c for c in ['name','TAD_id','RGG_id','B22_corr1_log10','Opposite_Charge_Number','base_prediction','final_hybrid_prediction','residual','absolute_error','overlap with 917 development set','model_manifest_sha256'] if c in nb.columns]]))
s7b=pd.read_csv(FIG5/'controls/figS7_final_model_controls/FigS7B_source.csv'); add_sheet(wb,'Shuffle perturbations',clean_df(s7b))
h=pd.read_csv(controls_pred/'histidine_pH6p5_final_predictions.csv'); add_sheet(wb,'Histidine pH6.5',clean_df(h[[c for c in ['name','B22_corr1_log10','Opposite_Charge_Number','base_prediction','final_hybrid_prediction','residual','absolute_error','model_manifest_sha256'] if c in h.columns]]))
basic=pd.read_csv(controls_pred/'basic_control_final_predictions_ALL80.csv'); add_sheet(wb,'Basic controls ALL80',clean_df(basic[[c for c in ['name','Corrected_B22_Hummer','B22_corr1_log10','target_evaluable','Opposite_Charge_Number','base_prediction','final_hybrid_prediction','residual','absolute_error','reason_target_undefined','model_manifest_sha256'] if c in basic.columns]]))
summ=[]; summ.append({'analysis':'non-basic WT','n':len(nb),**metrics(nb['B22_corr1_log10'],nb['final_hybrid_prediction']),'notes':'external/control WT pairs; overlap with 917 development set = '+str(int(nb['overlap with 917 development set'].sum()))})
for cls,label in [('TAD shuffle','TAD-shuffle perturbation'),('RGG shuffle','RGG-shuffle perturbation')]:
    sub=s7b[s7b['shuffle_class'].eq(cls)]; cm=corr(sub['delta_simulated'],sub['delta_predicted']); summ.append({'analysis':label,'n':cm['n'],'Pearson':cm['Pearson'],'Spearman':cm['Spearman'],'notes':'Delta convention is shuffle − WT; performance metrics are correlations only.'})
summ.append({'analysis':'histidine pH6.5','n':len(h),**metrics(h['B22_corr1_log10'],h['final_hybrid_prediction']),'notes':'Predictor features identical between pH7 and pH6.5; tests altered simulation condition.'})
be=basic[basic['target_evaluable'].astype(bool)]; summ.append({'analysis':'basic/noncanonical controls','n':len(be),**metrics(be['B22_corr1_log10'],be['final_hybrid_prediction']),'notes':'total n=80; evaluable n=78; positive-B22/outside-target-domain n=2; negative R² retained.'})
add_sheet(wb,'Summary',pd.DataFrame(summ)); add_sheet(wb,'Provenance',pd.DataFrame([{'item':'frozen model identity and cluster provenance','value':cluster_text},{'item':'final feature order','value':'; '.join(FEATURE_ORDER)},{'item':'source non-basic WT','value':str(controls_pred/'nonbasic_WT_final_predictions.csv')},{'item':'source shuffle perturbations','value':str(FIG5/'controls/figS7_final_model_controls/FigS7B_source.csv')},{'item':'source histidine pH6.5','value':str(controls_pred/'histidine_pH6p5_final_predictions.csv')},{'item':'source basic controls','value':str(controls_pred/'basic_control_final_predictions_ALL80.csv')}]))
path_s8=OUT/'Table_S8_Final_Model_Controls.xlsx'; save_wb(wb,path_s8)

(OUT/'FIG4_RESOURCE_TABLE_ADDITIONS.txt').write_text('Figure 4 supplementary resource table additions to keep separate from Tables S7 and S8:\n- WT vs shuffle simulation source data.\n- Histidine pH7 vs pH6.5 matched simulation source data.\n- RDF/convergence source data newly added to Figure 4 supplementary analyses.\nDo not rebuild the Figure 4 table in this task.\n')
lines=['FINAL TABLE S6-S8 AUDIT','']
for book in [path_s6,path_s7,path_s8]:
    wb2=load_workbook(book, read_only=True, data_only=True); lines.append(f'Workbook: {book}'); lines.append(f'Sheets: {wb2.sheetnames}')
    for ws in wb2.worksheets: lines.append(f'  {ws.title}: rows={max(ws.max_row-1,0) if ws.title!="README" else ws.max_row}')
    lines.append('')
lines += ['Verification summary:',f'Development OOF unique pairs: {pd.read_csv(fig5_final/"development_917_OOF_final_hybrid.csv")["pair_id"].nunique()} of 917 rows',f'Locked test unique pairs: {pd.read_csv(fig5_final/"test_1794_final_hybrid_predictions.csv")["pair_id"].nunique()} of 1794 rows',f'Full landscape rows: {len(pd.read_csv(fig5_final/"full_landscape_final_hybrid_predictions.csv", usecols=["pair_id"]))}',f'FINCHES S6A rows: {len(fa)}; S6B rows: {len(fb)}',f'Controls: non-basic WT={len(nb)}, shuffle={len(s7b)}, TAD shuffle={(s7b["shuffle_class"]=="TAD shuffle").sum()}, RGG shuffle={(s7b["shuffle_class"]=="RGG shuffle").sum()}, histidine pH6.5={len(h)}, basic total={len(basic)}, basic evaluable={int(basic["target_evaluable"].astype(bool).sum())}, positive-B22={int((~basic["target_evaluable"].astype(bool)).sum())}','Intentional omissions: historical B22_pred_hybrid, lambda=0 outputs, preliminary candidate predictions, Normalized OCN terminology, FINCHES/model-control mixing, and Figure 4 simulation-only controls.','Table S6 contains the definitive final hybrid-predictor resources underlying Figure 5 and Figure S5.','Table S7 contains the definitive FINCHES resources underlying Figure S6.','Table S8 contains the definitive frozen-model control/applicability resources underlying Figure S7.','Figure 4 simulation-only controls are intentionally excluded from Tables S7 and S8.']
(OUT/'FINAL_TABLE_S6_S8_AUDIT.txt').write_text('\n'.join(lines)+'\n')
print(path_s6); print(path_s7); print(path_s8)
