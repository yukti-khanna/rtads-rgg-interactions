#!/usr/bin/env python3
"""Build manuscript panels and tables from saved data plus the locked-test −NCP baseline check."""
from pathlib import Path
import argparse, importlib.util, json, hashlib, re
ROOT=Path(__file__).resolve().parents[1]
OUT=ROOT/'outputs/manuscript'

def module(name):
    p=ROOT/'scripts/manuscript_helpers'/name
    spec=importlib.util.spec_from_file_location(p.stem,p);m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m);return m

def label(value):
    return re.sub(r'(?<![\w])(?:AHD|TAD)(?![\w])','rAHD',value) if isinstance(value,str) else value

def sheet_name(kind, source_name):
    """Use tab labels that identify the supplementary table, not source files."""
    number={'predictor':8,'controls':9,'finches':10}[kind]
    control_names={
        'S8A_extended_source':'S9A Non-basic WT',
        'S8B_shuffle_source_audited':'S9B Shuffle controls',
        'S8D_extended_source':'S9C Basic controls',
        'S8B_absolute_metrics_original_new_combined':'S9D Absolute metrics',
        'S8B_delta_metrics_original_new_combined':'S9E Delta metrics',
        'predictor_delta_uncertainty_summary':'S9F Predictor uncertainty',
    }
    if kind=='controls' and source_name in control_names:
        return control_names[source_name]
    display={
        'readme':'overview',
        'development_oof_final_hybrid_bootstrap':'OOF bootstrap metrics',
        'development_oof_final_hybrid_predictions':'OOF pair predictions',
        'development_917_oof_final_hybrid':'development OOF',
        'test_1794_final_hybrid_predictions':'locked test 1794',
        'full_landscape_final_hybrid_predictions':'full landscape',
        'fig5b model development':'Fig4B model development',
        'fig5c patch ablation':'Fig4C patch ablation',
        'figs5 residual source':'FigS6 residual source',
        's6a finches simulation':'S10A FINCHES simulation',
        's6b finches landscape':'S10B FINCHES landscape',
    }.get(source_name.lower(), source_name.replace('_',' '))
    if display.startswith(f'S{number}'):
        return display[:31]
    prefix=f'S{number} '
    return (prefix+display)[:31]

def check():
    errors=[]
    manifest=json.loads((ROOT/'inputs/predictor/export_manifest.json').read_text())
    for item in manifest:
        p=ROOT/'inputs/predictor'/item['path']
        if not p.exists() or hashlib.sha256(p.read_bytes()).hexdigest()!=item['sha256']:errors.append(str(p))
    if errors:raise RuntimeError('Input checksum mismatch: '+str(errors))
    print('Predictor export checksums passed. No models loaded.')

def figures():
    import pandas as pd
    import matplotlib.pyplot as plt
    # Display-only substitution, leaving stored feature identifiers unchanged.
    from matplotlib.text import Text
    original=Text.set_text
    def set_text(self,s):return original(self,label(s))
    Text.set_text=set_text
    for name in ['Figure_4','Figure_S6','Figure_S7']:(OUT/name).mkdir(parents=True,exist_ok=True)
    m=module('update_final_figure_panels_from_predictions.py')
    m.make_main();m.make_s5()
    for old,new in [('Fig5D_development_OOF','Figure4B_development'),('Fig5E_heldout_1794','Figure4C_independent'),('Fig5F_full_landscape','Figure4D_landscape')]:
        for ext in ['pdf','svg','png']:(m.MAIN/f'{old}.{ext}').replace(m.MAIN/f'{new}.{ext}')
    # Keep only the accepted residual-correction panel in the manuscript set.
    for p in m.SUPP.glob('FigS5*'):
        if p.stem=='FigS5D_applied_residual_correction':p.replace(p.with_name('FigureS6A_residual_correction'+p.suffix))
        else:
            extra=OUT/'diagnostics/residual';extra.mkdir(parents=True,exist_ok=True);p.replace(extra/p.name)
    c=module('make_extended_s8_outputs.py')
    for src,stem in [('S8A_extended_source.csv','S8A_non_dataset_WT_extended'),('S8D_extended_source.csv','S8D_basic_controls_extended')]:
        df=pd.read_csv(ROOT/'inputs/controls'/src);valid=df.B22_corr1_log10.notna()
        c.plot_scatter(df[valid],stem,'' if valid.all() else f'{valid.sum()}/{len(df)} controls have B22 < 0')
        target='FigureS6D_nonbasic_WT' if 'S8A' in src else 'FigureS6E_basic_controls'
        for ext in ['pdf','svg','png']:(c.PANEL/f'{stem}.{ext}').replace(c.PANEL/f'{target}.{ext}')
    spec=importlib.util.spec_from_file_location('shuffle_export',ROOT/'scripts/export_updated_S6_shuffle_controls.py');b=importlib.util.module_from_spec(spec);spec.loader.exec_module(b)
    b.OUT=OUT/'Figure_S6'
    import shutil
    shutil.copy2(ROOT/'inputs/controls/S8B_shuffle_source_audited.csv',b.OUT/'S8B_shuffle_source_audited.csv')
    b.plot_absolute_b()
    u=module('run_predictor_shuffle_uncertainty.py')
    data=pd.read_csv(ROOT/'inputs/controls/predictor_shuffle_replica_data.csv');current=pd.read_csv(ROOT/'inputs/controls/current_delta_analysis.csv');summary=pd.read_csv(ROOT/'inputs/controls/predictor_delta_uncertainty_summary.csv')
    fig,axs=plt.subplots(1,2,figsize=(5.1,2.5))
    for ax,g in zip(axs,u.GROUPS):u.draw_panel(ax,g,'',data,current,summary)
    fig.subplots_adjust(left=.12,right=.98,bottom=.20,top=.87,wspace=.45)
    for ext in ['pdf','svg','png']:fig.savefig(b.OUT/f'FigureS6C_delta_uncertainty.{ext}',dpi=600)
    plt.close(fig)
    f=module('make_FigS6_FINCHES_scaled_panels.py')
    for src,fun in [('S6A FINCHES simulation.csv',f.save_panel_a),('S6B FINCHES landscape.csv',f.save_panel_b)]:
        df=pd.read_csv(ROOT/'inputs/finches'/src).rename(columns={'scaled FINCHES value':'eps_scaled','scaled interaction value':'y_scaled','scaled final prediction':'y_scaled'})
        fun(df,f.pearson(df.eps_scaled,df.y_scaled))
    for p in f.OUTDIR.glob('FigS6*'):p.replace(p.with_name(p.name.replace('FigS6','FigureS7')))
    Text.set_text=original

def tables():
    # User-run portable workbook writer; no model calculation.
    import pandas as pd
    from openpyxl import Workbook
    from openpyxl.styles import Font
    target=OUT/'tables';target.mkdir(parents=True,exist_ok=True)
    for kind,num in [('predictor',8),('controls',9),('finches',10)]:
        w=Workbook();w.remove(w.active)
        note=w.create_sheet(f'S{num} Overview');note.append(['Item','Value']);note.append(['Workbook',f'Table S{num}: '+{'predictor':'final predictor resources','controls':'frozen-model controls','finches':'FINCHES resources'}[kind]]);note.append(['Terminology','rAHD; internal feature identifiers preserved']);note.append(['Controls caveat','Stored accepted predictions; global residue-table discrepancy remains unresolved']);note.append(['Histidine','Retained separately; excluded from these tables'])
        files=sorted((ROOT/'inputs'/kind).glob('*.csv'))
        if kind=='controls':
            names=['S8A_extended_source.csv','S8B_shuffle_source_audited.csv','S8D_extended_source.csv','S8B_absolute_metrics_original_new_combined.csv','S8B_delta_metrics_original_new_combined.csv','predictor_delta_uncertainty_summary.csv']
            files=[ROOT/'inputs/controls'/n for n in names]
        for p in files:
            if 'histidine' in p.name.lower() or 'S8C' in p.name:continue
            df=pd.read_csv(p)
            if p.name.lower() == 'readme.csv':
                # Each exported workbook gets the manuscript table number
                # assigned by this build, regardless of its input README label.
                df=df.replace(to_replace=r'Table\s+S\d+',value=f'Table S{num}',regex=True)
            if kind=='controls' and p.name=='S8B_shuffle_source_audited.csv':
                features=pd.read_csv(ROOT/'inputs/controls/S8B_extended_source.csv')
                extra=[c for c in features.columns if c not in df.columns and c!='name']
                df=df.merge(features[['name']+extra],left_on='shuffle_pair_id',right_on='name',how='left',validate='one_to_one')
            name=sheet_name(kind,p.stem)
            if name in w.sheetnames:raise ValueError(f'Duplicate sheet: {name}')
            s=w.create_sheet(name);s.append([label(x) for x in df.columns])
            for row in df.itertuples(index=False,name=None):s.append([None if pd.isna(x) else label(x) for x in row])
            s.freeze_panes='A2';s.auto_filter.ref=s.dimensions
            for cell in s[1]:cell.font=Font(bold=True)
        w.save(target/f'Table_S{num}_{kind}.xlsx')
        from clean_supplementary_workbook import clean_workbook
        clean_workbook(target/f'Table_S{num}_{kind}.xlsx')
    # Histidine is retained separately, never added to a manuscript table.
    hist=ROOT/'inputs/controls/histidine_pH6p5_final_predictions.csv'
    if hist.exists():
        import shutil
        dest=OUT/'analyses/histidine';dest.mkdir(parents=True,exist_ok=True);shutil.copy2(hist,dest/hist.name)

def performance_by_minus_ncp_bin():
    path=ROOT/'scripts/compare_performance_by_minus_NCP_bin.py'
    spec=importlib.util.spec_from_file_location('performance_by_minus_ncp_bin',path)
    analysis=importlib.util.module_from_spec(spec);spec.loader.exec_module(analysis)
    analysis.build()

def heldout_ncp_model_rmse_by_bin():
    path=ROOT/'scripts/compare_heldout_ncp_only_vs_full13_by_bin.py'
    spec=importlib.util.spec_from_file_location('heldout_ncp_model_rmse_by_bin',path)
    analysis=importlib.util.module_from_spec(spec);spec.loader.exec_module(analysis)
    analysis.main()

def main():
    p=argparse.ArgumentParser();p.add_argument('--check-only',action='store_true');p.add_argument('--tables-only',action='store_true');a=p.parse_args();check()
    if a.tables_only:
        tables()
    elif not a.check_only:
        figures();tables();performance_by_minus_ncp_bin();heldout_ncp_model_rmse_by_bin()
        expected=[OUT/'Figure_4/Figure4D_landscape.pdf',OUT/'Figure_S6/FigureS6A_residual_correction.pdf',OUT/'Figure_S6/FigureS6F_performance_by_minus_NCP_bin.pdf',OUT/'Figure_S6/FigureS6G_heldout_ncp_model_rmse_by_bin.pdf',OUT/'Figure_S6/FigureS6G_heldout_ncp_model_rmse_by_bin.svg',OUT/'Figure_S6/FigureS6G_heldout_ncp_model_rmse_by_bin.png',OUT/'Figure_S7',OUT/'tables/Table_S8_predictor.xlsx',OUT/'tables/Table_S9_controls.xlsx',OUT/'tables/Table_S10_finches.xlsx',OUT/'performance_by_minus_NCP_bin/performance_by_minus_NCP_bin.csv',OUT/'performance_by_minus_NCP_bin/sampling_representation_by_minus_NCP_bin.csv',OUT/'performance_by_minus_NCP_bin/cohort_summary.csv',OUT/'performance_by_minus_NCP_bin/performance_by_minus_NCP_bin.svg',OUT/'performance_by_minus_NCP_bin/performance_by_minus_NCP_bin.png']
        missing=[str(p) for p in expected if not p.exists()]
        if missing:raise RuntimeError('Manuscript build is incomplete; missing outputs: '+', '.join(missing))
        from openpyxl import load_workbook
        for number,kind in [(8,'predictor'),(9,'controls'),(10,'finches')]:
            book=load_workbook(OUT/'tables'/f'Table_S{number}_{kind}.xlsx',read_only=True,data_only=True)
            label_text=' '.join(str(c.value or '') for row in book[f'S{number} Overview'].iter_rows() for c in row)
            if f'Table S{number}' not in label_text:
                raise RuntimeError(f'Table S{number} ({kind}) workbook label is incorrect.')
        print('Verified manuscript figures and Tables S8 (predictor), S9 (controls), and S10 (FINCHES).')
if __name__=='__main__':main()
