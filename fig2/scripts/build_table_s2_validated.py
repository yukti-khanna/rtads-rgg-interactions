#!/usr/bin/env python3
"""Build a validated Table S2 from the authoritative Figure 1 NMR sources.

This script deliberately does not use the historical qualitative peptide-call
sheet as the source for Figure 1A.  Figure 1A is reconstructed from the full
per-residue CSP workbook; Figure 1B and the shuffle analyses remain tied to the
NMR parent master table.  Historical sheets are retained and clearly labelled
for auditability.
"""
from __future__ import annotations

import argparse
import json
import re
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


PROTEINS = ["FOXM1", "p53", "FOXO4", "LEF1"]
PEPTIDES = ["RRGG", "KRGG", "LRGG", "CRGG", "ERGG"]
BLOCKS = {
    "FOXM1": (1, {"RRGG": 2, "LRGG": 3, "KRGG": 4}),
    "p53": (8, {"CRGG": 9, "LRGG": 10, "KRGG": 11, "RRGG": 12}),
    "FOXO4": (15, {"CRGG": 16, "LRGG": 17, "KRGG": 18, "RRGG": 19}),
    "LEF1": (24, {"CRGG": 25, "KRGG": 26, "RRGG": 27}),
}
FOXM1_EQ = {"RRGG": .58, "KRGG": 1.0, "LRGG": 1.0}


def ra(value):
    if not isinstance(value, str):
        return value
    value = re.sub(r"\bAHDs\b", "rAHDs", value)
    return re.sub(r"\bAHD\b", "rAHD", value)


def csp_summary(path: Path) -> tuple[pd.DataFrame, pd.DataFrame]:
    raw = pd.read_excel(path, sheet_name="foxm1", header=None)
    rows = []
    for protein in PROTEINS:
        residue_col, peptide_cols = BLOCKS[protein]
        valid = raw.iloc[:, residue_col].astype(str).str.startswith("@")
        for peptide in PEPTIDES:
            if peptide not in peptide_cols:
                rows.append({"protein": protein, "peptide": peptide, "status": "not available"})
                continue
            values = pd.to_numeric(raw.loc[valid, peptide_cols[peptide]], errors="coerce").dropna()
            measured = float(values.mean())
            eq = FOXM1_EQ.get(peptide, np.nan) if protein == "FOXM1" else np.nan
            plotted = measured * .58 / eq if np.isfinite(eq) else measured
            rows.append({
                "protein": protein, "peptide": peptide, "status": "measured",
                "n_residue_peaks": len(values), "mean_CSP_measured": measured,
                "SD_CSP_measured": float(values.std(ddof=1)),
                "median_CSP_measured": float(values.median()),
                "ligand_equivalents_measured": eq,
                "mean_CSP_used_for_Figure1A": plotted,
            })
    summary = pd.DataFrame(rows)
    summary["relative_CSP_within_TF"] = np.nan
    for protein, idx in summary.groupby("protein").groups.items():
        maximum = summary.loc[idx, "mean_CSP_used_for_Figure1A"].max()
        summary.loc[idx, "relative_CSP_within_TF"] = summary.loc[idx, "mean_CSP_used_for_Figure1A"] / maximum
    return raw, summary


def replace_sheet(workbook, name: str, data: pd.DataFrame, note: str):
    if name in workbook.sheetnames:
        del workbook[name]
    sheet = workbook.create_sheet(name)
    sheet.append([note])
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(1, len(data.columns)))
    sheet["A1"].font = Font(bold=True)
    sheet["A1"].fill = PatternFill("solid", fgColor="FFF2CC")
    sheet.append(list(data.columns))
    for cell in sheet[2]:
        cell.font = Font(bold=True)
    for row in data.itertuples(index=False, name=None):
        sheet.append([
            None if (isinstance(value, (float, np.floating)) and not np.isfinite(value)) else value
            for value in row
        ])
    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = sheet.dimensions
    for column_index, column in enumerate(sheet.columns, start=1):
        width = min(45, max(10, max(len(str(cell.value or "")) for cell in column) + 2))
        # The source-note row is merged, so its non-anchor cells are
        # MergedCell objects without ``column_letter``. Use the stable numeric
        # position instead of inspecting the first cell in each column.
        sheet.column_dimensions[get_column_letter(column_index)].width = width


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    root = Path(__file__).resolve().parents[1]
    parser.add_argument("--raw-csp", type=Path, default=root / "inputs/tables/xrgg csps.xlsx")
    parser.add_argument("--nmr-master", type=Path, default=root / "inputs/tables/xrgg_AHD_parent_master.csv")
    parser.add_argument("--template", type=Path, default=root / "final_manuscript_pipeline/inputs/curated_tables/Table_S2_Figure1_NMR_CURATED.xlsx")
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--audit-json", type=Path)
    parser.add_argument("--overwrite", action="store_true")
    args = parser.parse_args()
    for path in (args.raw_csp, args.nmr_master, args.template):
        if not path.is_file():
            raise FileNotFoundError(path)
    if args.output.exists() and not args.overwrite:
        raise FileExistsError(f"Refusing to overwrite {args.output}")

    raw, summary = csp_summary(args.raw_csp)
    master = pd.read_csv(args.nmr_master)
    master = master.apply(lambda column: column.map(ra) if column.dtype == object else column)
    workbook = load_workbook(args.template)
    renames = {
        "S2A_Peptide_screen_DRAFT": "S2A_Historical_calls",
        "S2G_FOXM1_CSP_raw": "S2G_Historical_CSP_subset",
    }
    for old, new in renames.items():
        if old in workbook.sheetnames:
            workbook[old].title = new
    for sheet in workbook.worksheets:
        for row in sheet.iter_rows():
            for cell in row:
                cell.value = ra(cell.value)

    replace_sheet(
        workbook, "S2A_CSP_matrix_FINAL", summary,
        "AUTHORITATIVE Figure 1A source summary: means from the full per-residue xrgg csps.xlsx workbook. FOXM1 KRGG/LRGG means are scaled from 1.00 to 0.58 ligand equivalents for plotting; FOXM1 RRGG was measured at 0.58 equivalents.",
    )
    raw.columns = [f"source_column_{i}" for i in range(raw.shape[1])]
    replace_sheet(
        workbook, "S2G_CSP_raw_EXACT", raw,
        "Exact foxm1 worksheet imported without row filtering from xrgg csps.xlsx; blank cells are preserved as blanks/NaN by the spreadsheet reader.",
    )
    replace_sheet(
        workbook, "S2H_NMR_master_EXACT", master,
        "Exact rows from xrgg_AHD_parent_master.csv (display terminology updated only from AHD to rAHD where applicable).",
    )
    readme = workbook["README"] if "README" in workbook.sheetnames else workbook.create_sheet("README", 0)
    readme.insert_rows(1, 8)
    statements = [
        "SOURCE-OF-TRUTH NOTICE",
        "Figure 1A quantitative CSP matrix: S2A_CSP_matrix_FINAL, derived from inputs/tables/xrgg csps.xlsx.",
        "The historical qualitative-call sheet and historical FOXM1 subset are retained for audit only and must not be used to regenerate Figure 1A.",
        "Figure 1B and WT/shuffle NMR analyses: inputs/tables/xrgg_AHD_parent_master.csv (copied into S2H_NMR_master_EXACT).",
        "Missing TF-peptide combinations remain 'not available'; no qualitative calls are inferred from absent quantitative data.",
        "FOXM1 plotting normalization: RRGG measured at 0.58 eq; KRGG/LRGG measured at 1.00 eq and linearly scaled to 0.58 eq.",
        "Displayed terminology uses rAHD; identifiers and numerical measurements are unchanged.",
        "",
    ]
    for index, value in enumerate(statements, 1):
        readme.cell(index, 1, value)
        if index == 1:
            readme.cell(index, 1).font = Font(bold=True)

    args.output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(args.output)
    check = load_workbook(args.output, read_only=True, data_only=False)
    check.close()
    audit = {
        "figure1A_authoritative_source": str(args.raw_csp.resolve()),
        "figure1B_and_shuffle_authoritative_source": str(args.nmr_master.resolve()),
        "template_retained_for_historical_audit": str(args.template.resolve()),
        "measured_TF_peptide_combinations": int(summary.status.eq("measured").sum()),
        "figure1A_rule": "per-residue CSP mean; FOXM1 KRGG/LRGG scaled from 1.00 to 0.58 eq",
        "historical_sheets_are_authoritative": False,
        "figure1A_summary": [
            {key: (None if pd.isna(value) else value) for key, value in record.items()}
            for record in summary.to_dict(orient="records")
        ],
        "output": str(args.output.resolve()),
    }
    audit_path = args.audit_json or args.output.with_suffix(".audit.json")
    audit_path.parent.mkdir(parents=True, exist_ok=True)
    audit_path.write_text(json.dumps(audit, indent=2) + "\n")
    print(json.dumps(audit, indent=2))


if __name__ == "__main__":
    main()
