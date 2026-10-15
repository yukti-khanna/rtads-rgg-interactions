#!/usr/bin/env python3
"""Build the audited Figure 1/S2 and Table S2-S4 manuscript package.

Analysis mode is explicit: ``rerun`` rebuilds the auditable NMR/CSP summary and
Figure 1C cluster bootstrap; ``reuse`` consumes a previously validated analysis
cache.  BLAST and external-AD recurrence remain separate opt-in analyses.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
from pathlib import Path


ROOT = Path(__file__).resolve().parents[1]
PIPELINE = ROOT / "final_manuscript_pipeline"
CURATED = PIPELINE / "inputs" / "curated_tables"
REFERENCES = PIPELINE / "inputs" / "reference_figures"


def run(command: list[str], cwd: Path) -> None:
    print("RUN:", " ".join(map(str, command)))
    subprocess.run(command, cwd=cwd, check=True)


def copy_file(source: Path, destination: Path) -> Path:
    if not source.is_file():
        raise FileNotFoundError(source)
    destination.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source, destination)
    return destination


def replace_rahd_labels(value: str) -> str:
    """Update displayed terminology without changing identifiers or numbers."""
    value = re.sub(r"\bAHDs\b", "rAHDs", value)
    value = re.sub(r"\bAHD\b", "rAHD", value)
    return value


def write_curated_tables(output: Path, validated_table_s2: Path) -> list[Path]:
    """Create submission copies of the frozen curated Tables S2-S4."""
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        raise RuntimeError("openpyxl is required to create Tables S2-S4") from exc

    mapping = {
        "Table_S3_Figure1C_Features_AUDIT.xlsx": "Table_S3_Figure1C_Features_FINAL.xlsx",
        "Table_S4_rAHD_Classifier_Benchmark_BLAST_CURATED.xlsx":
            "Table_S4_rAHD_Classifier_Benchmark_BLAST_FINAL.xlsx",
    }
    generated: list[Path] = []
    generated.append(copy_file(
        validated_table_s2,
        output / "supplementary_tables/Table_S2_Figure1_NMR_FINAL.xlsx",
    ))
    for source_name, final_name in mapping.items():
        source = CURATED / source_name
        if not source.is_file():
            raise FileNotFoundError(source)
        workbook = load_workbook(source)
        changed = 0
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows():
                for cell in row:
                    if isinstance(cell.value, str):
                        revised = replace_rahd_labels(cell.value)
                        if revised != cell.value:
                            cell.value = revised
                            changed += 1
        destination = output / "supplementary_tables" / final_name
        destination.parent.mkdir(parents=True, exist_ok=True)
        workbook.save(destination)
        # Reopen to ensure the submission workbook is structurally readable.
        checked = load_workbook(destination, read_only=True, data_only=False)
        if not checked.sheetnames:
            raise ValueError(f"No worksheets in {destination}")
        checked.close()
        print(f"Created {destination} ({changed} AHD -> rAHD label updates)")
        generated.append(destination)
    return generated


def copy_formats(source_stem: Path, destination_stem: Path,
                 formats: tuple[str, ...] = ("pdf", "png", "svg")) -> list[Path]:
    generated = []
    for extension in formats:
        source = source_stem.with_suffix("." + extension)
        if source.is_file():
            generated.append(copy_file(source, destination_stem.with_suffix("." + extension)))
    if not generated:
        raise FileNotFoundError(f"No generated formats found for {source_stem}")
    return generated


def assemble_png_panels(paths: list[Path], output_stem: Path, layout: tuple[int, int],
                        letters: list[str], frozen_crop: tuple[Path, tuple[float, float, float, float]] | None = None) -> list[Path]:
    """Make a simple manuscript assembly; individual vector panels remain primary."""
    import matplotlib as mpl
    import matplotlib.pyplot as plt

    mpl.rcParams.update({"font.family": "sans-serif", "font.sans-serif": ["Arial", "Helvetica", "DejaVu Sans"],
                         "pdf.fonttype": 42, "ps.fonttype": 42, "svg.fonttype": "none"})
    rows, columns = layout
    fig, axes = plt.subplots(rows, columns, figsize=(7.2, 3.7 * rows), squeeze=False)
    panel_images = []
    if frozen_crop is not None:
        source, (left, top, right, bottom) = frozen_crop
        image = plt.imread(source)
        height, width = image.shape[:2]
        panel_images.append(image[int(top*height):int(bottom*height), int(left*width):int(right*width)])
    panel_images.extend(plt.imread(path) for path in paths)
    for index, ax in enumerate(axes.flat):
        ax.axis("off")
        if index >= len(panel_images):
            continue
        ax.imshow(panel_images[index])
        if index < len(letters) and letters[index]:
            ax.text(0.005, 0.995, letters[index], transform=ax.transAxes,
                    ha="left", va="top", fontsize=9, fontweight="bold")
    fig.subplots_adjust(left=0.01, right=0.99, bottom=0.01, top=0.99, wspace=0.04, hspace=0.06)
    output_stem.parent.mkdir(parents=True, exist_ok=True)
    outputs = []
    for extension in ("pdf", "svg", "png"):
        path = output_stem.with_suffix("." + extension)
        fig.savefig(path, dpi=600 if extension == "png" else None, facecolor="white", bbox_inches="tight")
        outputs.append(path)
    plt.close(fig)
    return outputs


def prepare_analysis(args: argparse.Namespace, output: Path) -> dict[str, Path]:
    """Rerun or reuse the manuscript analysis cache without training models."""
    py = sys.executable
    analysis = output / "analysis"
    table_s2 = analysis / "Table_S2_Figure1_NMR_VALIDATED.xlsx"
    fig1c_stats = analysis / "Fig1C_cluster_bootstrap_statistics.csv"
    sequence_source = PIPELINE / "inputs/analysis_cache/sequence_similarity/FigS3_sequence_similarity_source.csv"
    base_revision_config = PIPELINE / "inputs/final_pipeline_inputs.json"
    if args.analysis_mode == "rerun":
        analysis.mkdir(parents=True, exist_ok=True)
        overwrite = ["--overwrite"] if args.overwrite_analysis else []
        if fig1c_stats.exists() and not args.overwrite_analysis:
            raise FileExistsError(f"Analysis output exists: {fig1c_stats}; use --overwrite-analysis or --analysis-mode reuse")
        run([py, "scripts/fig1C_cluster_bootstrap.py", "--project-root", str(ROOT),
             "--output", str(fig1c_stats), *overwrite], ROOT)
        run([py, "scripts/build_table_s2_validated.py", "--output", str(table_s2),
             "--audit-json", str(analysis / "Table_S2_source_audit.json"), *overwrite], ROOT)
    else:
        missing = [path for path in (table_s2, fig1c_stats) if not path.is_file()]
        if missing:
            raise FileNotFoundError(
                "Reuse mode requires a validated cache. Missing:\n" +
                "\n".join(map(str, missing)) +
                "\nRun once with --analysis-mode rerun."
            )

    if args.rerun_blast:
        blast_dir = analysis / "sequence_similarity"
        starting = PIPELINE / "inputs/analysis_cache/sequence_similarity/original_9_starting_regions_no_AXIN1.fasta"
        run([py, "scripts/merged_fragments_vs_9seeds_blast.py",
             "--query-fasta", str(ROOT / "inputs/fasta/100_all_at2_frags_dis.fasta"),
             "--starting-fasta", str(starting), "--outdir", str(blast_dir),
             "--threads", str(args.threads), "--run-blast", "--overwrite"], ROOT)
        run([py, "scripts/finalize_figS3_sequence_similarity.py", "--analysis-dir", str(blast_dir),
             "--query-fasta", str(ROOT / "inputs/fasta/100_all_at2_frags_dis.fasta"),
             "--threads", str(args.threads), "--overwrite"], ROOT)
        sequence_source = blast_dir / "FigS3_sequence_similarity_source.csv"

    if args.rerun_external_ad:
        required = {
            "--candidate-window-fasta": args.candidate_window_fasta,
            "--window-votes": args.window_votes,
            "--candidate-region-table": args.candidate_region_table,
        }
        missing = [flag for flag, path in required.items() if path is None]
        if missing:
            raise ValueError("--rerun-external-ad requires " + ", ".join(missing))
        run([py, "scripts/external_AD_ensemble_recurrence.py",
             "--project-root", str(ROOT),
             "--window-votes", str(args.window_votes.resolve()),
             "--candidate-window-fasta", str(args.candidate_window_fasta.resolve()),
             "--candidate-region-table", str(args.candidate_region_table.resolve()),
             "--outdir", str(analysis / "external_AD_ensemble_recurrence"),
             "--final-consensus-threshold", "1.0"], ROOT)
        config = json.loads(base_revision_config.read_text())
        recurrence = analysis / "external_AD_ensemble_recurrence"
        config.update({
            "recurrence_curves": str(recurrence / "external_AD_recovery_by_consensus_threshold.csv"),
            "external_recurrence": str(recurrence / "external_AD_window_recurrence.csv"),
            "matched_controls": str(recurrence / "external_AD_matched_control_pairs.csv"),
            "control_statistics": str(recurrence / "external_AD_vs_matched_control_recurrence.csv"),
            "nonfamily_curves": str(recurrence / "external_AD_recurrence_nonFOX_sensitivity.csv"),
        })
    else:
        config = json.loads(base_revision_config.read_text())
    config["output_directory"] = str(output / "plot_cache")
    revision_config = analysis / "figure2_revision_inputs_selected.json"
    revision_config.write_text(json.dumps(config, indent=2) + "\n")
    return {"table_s2": table_s2, "fig1c_stats": fig1c_stats,
            "sequence_source": sequence_source, "revision_config": revision_config}


def build_figures(output: Path, cache: dict[str, Path]) -> list[Path]:
    py = sys.executable
    generated: list[Path] = []

    # Figure 1A/B and Figure S2B from the curated NMR master inputs.
    run([py, "scripts/make_fig1A_xrgg_csp_matrix.py"], ROOT)
    run([py, "xrgg_panel_scripts/build_xrgg_ahd_panels.py", "--project-root", str(ROOT)], ROOT)

    fig1 = output / "figures" / "Figure1" / "panels"
    generated += copy_formats(ROOT / "outputs/fig1/Fig1A_XRGG_CSP_matrix", fig1 / "Figure1A")
    generated += copy_formats(ROOT / "xrgg_panels/Fig_main_RRGG_reference_Kd_B", fig1 / "Figure1B", ("pdf", "png"))

    # Figure 1C uses raw values, unchanged g point estimates, and cluster CIs.
    fig1c_dir = fig1 / "Figure1C_build"
    run([py, "scripts/remake_fig1C_boxplots.py", "--project-root", str(ROOT),
         "--statistics", str(cache["fig1c_stats"]),
         "--output-dir", str(fig1c_dir), "--overwrite"], ROOT)
    generated += copy_formats(fig1c_dir / "Fig1C_boxplots_corrected", fig1 / "Figure1C")
    generated.append(copy_file(fig1c_dir / "Fig1C_boxplots_source_data.csv",
                               fig1 / "Figure1C_source_data.csv"))

    # Cached analytical classifier panels only; no schematic, scoring or training.
    run([py, "scripts/plot_figure1_classifier_panels.py", "--project-root", str(ROOT),
         "--config", str(cache["revision_config"])], ROOT)
    config = json.loads(cache["revision_config"].read_text())
    manuscript_source = ROOT / config["output_directory"] / "manuscript"
    generated += copy_formats(manuscript_source / "Figure1/Figure1E_external_AD_recurrence",
                              fig1 / "Figure1E")

    # Supplementary Figure S2.
    figs2 = output / "figures" / "FigureS2" / "panels"
    generated += copy_formats(ROOT / "xrgg_panels/FigS_RRGG_WT_vs_shuffle",
                              figs2 / "FigureS2B", ("pdf", "png"))
    run([py, "scripts/plot_FigS3_sequence_similarity_FINAL.py",
         "--source", str(cache["sequence_source"]),
         "--outdir", str(figs2), "--overwrite"], ROOT)
    for extension in ("pdf", "png"):
        generated.append(copy_file(figs2 / f"FigS3_sequence_similarity_FINAL.{extension}",
                                   figs2 / f"FigureS2C.{extension}"))
    for letter, source_name in {
        "D": "FigureS2D_ensemble_support_ECDF",
        "E": "FigureS2E_selected_support_thresholds",
        "F": "FigureS2F_nonfamily_sensitivity",
    }.items():
        generated += copy_formats(manuscript_source / f"FigureS2/{source_name}",
                                  figs2 / f"FigureS2{letter}")
    # The original editable S2A NMR assembly was not supplied. Preserve it as a
    # frozen crop from the submitted supplementary figure; all other panels are regenerated.
    s2_paths = [figs2 / f"FigureS2{letter}.png" for letter in "BCDEF"]
    generated += assemble_png_panels(
        s2_paths, output / "figures/FigureS2/FigureS2_final", (3, 2),
        ["A", "B", "C", "", "", ""],
        frozen_crop=(REFERENCES / "FigureS2_manuscript_reference.png", (0.0, 0.0, 1.0, 0.245)),
    )

    return generated


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def write_manifest(output: Path) -> Path:
    files = []
    for path in sorted(p for p in output.rglob("*") if p.is_file()):
        if path.name == "manifest.json":
            continue
        files.append({"path": str(path.relative_to(output)), "bytes": path.stat().st_size,
                      "sha256": sha256(path)})
    manifest = {
        "scope": ["Figure 1 analytical panels", "Figure S2 analytical panels", "Table S2", "Table S3", "Table S4"],
        "explicitly_excluded": ["Figure 1D workflow schematic", "Figure 2", "Figure S3"],
        "terminology": "rAHD = regulatory acidic-hydrophobic disordered region",
        "analysis_policy": "no model training or candidate rescoring; NMR/CSP and cluster-bootstrap analyses may be rerun explicitly; BLAST/external-AD recurrence are opt-in",
        "files": files,
    }
    path = output / "manifest.json"
    path.write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    return path


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output-dir", type=Path,
                        default=ROOT / "outputs/final_manuscript_package")
    parser.add_argument("--figures-only", action="store_true")
    parser.add_argument("--tables-only", action="store_true")
    parser.add_argument("--analysis-mode", choices=("rerun", "reuse"), default="rerun",
                        help="rerun validated NMR/CSP and Figure 1C analyses, or reuse this output directory's cache")
    parser.add_argument("--overwrite-analysis", action="store_true")
    parser.add_argument("--rerun-blast", action="store_true",
                        help="Opt in to rerunning the 485-region versus nine-seed BLAST audit")
    parser.add_argument("--rerun-external-ad", action="store_true",
                        help="Opt in to rerunning external-AD ensemble recurrence")
    parser.add_argument("--candidate-window-fasta", type=Path,
                        help="Complete prediction-universe 15-aa FASTA; required with --rerun-external-ad")
    parser.add_argument("--window-votes", type=Path,
                        help="Recovered 168-model candidate vote table; required with --rerun-external-ad")
    parser.add_argument("--candidate-region-table", type=Path,
                        help="Complete biological candidate-region table; required with --rerun-external-ad")
    parser.add_argument("--threads", type=int, default=8)
    args = parser.parse_args()
    if args.figures_only and args.tables_only:
        parser.error("Choose at most one of --figures-only and --tables-only")
    output = args.output_dir.resolve()
    output.mkdir(parents=True, exist_ok=True)
    cache = prepare_analysis(args, output)
    if not args.tables_only:
        build_figures(output, cache)
    if not args.figures_only:
        write_curated_tables(output, cache["table_s2"])
    manifest = write_manifest(output)
    print(f"Final manuscript package: {output}")
    print(f"Manifest: {manifest}")


if __name__ == "__main__":
    main()
